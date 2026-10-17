"""
Genera el reporte de resultados de pruebas de Conexión360 en Excel y HTML.

Fuentes:
  - casos_azure.json: los test cases de Azure DevOps (catálogo leído por MCP).
  - ../katalon/Conexion360/Reports/**/JUnit_Report.xml: ejecuciones de las suites de Katalon.
  - resultados_manuales.csv: resultados de los casos manuales y parciales (se llena a mano).

Uso (desde esta carpeta):  py generar_reporte.py
Salida: Reporte_Resultados_<fecha>.xlsx y Reporte_Resultados_<fecha>.html en esta carpeta.
"""
import csv
import html
import json
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
PROYECTO_KATALON = AQUI.parent / 'katalon' / 'Conexion360'
AZURE = 'https://dev.azure.com/ingenieria-requisitos/Conexi%C3%B3n360/_workitems/edit/'

RESULTADOS = ['Pasó', 'Falló', 'Error', 'Bloqueado', 'No ejecutado', 'Pendiente manual']
ESTADO_KATALON = {'PASSED': 'Pasó', 'FAILED': 'Falló', 'ERROR': 'Error', 'SKIPPED': 'No ejecutado',
                  'INCOMPLETE': 'Error'}
COLOR = {  # relleno Excel / color HTML por resultado
    'Pasó': 'D9F2E3', 'Falló': 'FAD7D4', 'Error': 'F8E1C4', 'Bloqueado': 'E8E0F5',
    'No ejecutado': 'EDEFF2', 'Pendiente manual': 'FFF2CC',
}


# ---------------------------------------------------------------- datos
def cargar_casos():
    return json.loads((AQUI / 'casos_azure.json').read_text(encoding='utf-8'))


def cargar_hallazgos():
    archivo = AQUI / 'hallazgos.json'
    if not archivo.exists():
        return {'fecha': '', 'metodo': '', 'hallazgos': []}
    return json.loads(archivo.read_text(encoding='utf-8'))


def casos_katalon():
    """ID de Azure -> ruta del test case en Katalon (por el prefijo TC<id>_ del nombre)."""
    mapa = {}
    for tc in (PROYECTO_KATALON / 'Test Cases').rglob('TC*.tc'):
        m = re.match(r'TC(\d+)_', tc.name)
        if m:
            mapa[int(m.group(1))] = 'Test Cases/' + tc.relative_to(PROYECTO_KATALON / 'Test Cases').with_suffix('').as_posix()
    return mapa


def leer_junit():
    """Todas las ejecuciones de Katalon encontradas en Reports/, de la más antigua a la más reciente."""
    ejecuciones = []
    for archivo in (PROYECTO_KATALON / 'Reports').rglob('JUnit_Report.xml'):
        try:
            raiz = ET.parse(archivo).getroot()
        except ET.ParseError:
            continue
        for suite in raiz.iter('testsuite'):
            fecha = suite.get('timestamp') or datetime.fromtimestamp(archivo.stat().st_mtime).strftime('%Y-%m-%d %H:%M:%S')
            for tc in suite.iter('testcase'):
                nombre = tc.get('name', '')
                m = re.search(r'TC(\d+)_', nombre)
                if not m:
                    continue
                estado = (tc.get('status') or '').upper()
                problema = tc.find('failure') if tc.find('failure') is not None else tc.find('error')
                if not estado:
                    estado = 'FAILED' if tc.find('failure') is not None else 'ERROR' if tc.find('error') is not None \
                        else 'SKIPPED' if tc.find('skipped') is not None else 'PASSED'
                mensaje = ''
                if problema is not None:
                    mensaje = (problema.get('message') or problema.text or '').strip()
                ejecuciones.append({
                    'id': int(m.group(1)), 'fecha': fecha.replace('T', ' ')[:19], 'suite': suite.get('name', ''),
                    'caso': nombre, 'resultado': ESTADO_KATALON.get(estado, 'Error'),
                    'duracion': round(float(tc.get('time') or 0), 1), 'mensaje': mensaje[:1000],
                    'reporte': archivo.relative_to(PROYECTO_KATALON).as_posix(),
                })
    return sorted(ejecuciones, key=lambda e: e['fecha'])


def leer_manuales(casos):
    """Resultados de casos manuales y parciales. Si el CSV no existe, crea la plantilla."""
    archivo = AQUI / 'resultados_manuales.csv'
    if not archivo.exists():
        with archivo.open('w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['ID', 'Titulo', 'Resultado', 'Fecha', 'Ejecutor', 'Observaciones'])
            for c in casos:
                if c['tipo'] != 'Automatizado':
                    w.writerow([c['id'], c['titulo'], '', '', '', ''])
    manuales = {}
    with archivo.open(encoding='utf-8-sig', newline='') as f:
        for fila in csv.DictReader(f, delimiter=';'):
            resultado = (fila.get('Resultado') or '').strip()
            if fila.get('ID', '').strip().isdigit() and resultado:
                manuales[int(fila['ID'])] = {
                    'resultado': resultado if resultado in RESULTADOS else 'Error',
                    'fecha': (fila.get('Fecha') or '').strip(), 'mensaje': (fila.get('Observaciones') or '').strip(),
                    'ejecutor': (fila.get('Ejecutor') or '').strip(),
                }
    return manuales


def consolidar():
    datos = cargar_casos()
    casos, katalon, ejecuciones = datos['casos'], casos_katalon(), leer_junit()
    manuales = leer_manuales(casos)
    por_caso = {}
    for h in cargar_hallazgos()['hallazgos']:
        for i in h['casos']:
            por_caso.setdefault(i, []).append(h['id'])
    ultima = {}
    for e in ejecuciones:  # la más reciente gana
        ultima[e['id']] = e
    filas = []
    for c in casos:
        fila = dict(c, katalon=katalon.get(c['id'], ''), resultado='No ejecutado', duracion=None,
                    fecha='', fuente='', mensaje='', hallazgos=', '.join(por_caso.get(c['id'], [])))
        if c['id'] in ultima:
            e = ultima[c['id']]
            fila.update(resultado=e['resultado'], duracion=e['duracion'], fecha=e['fecha'],
                        fuente='Katalon (' + e['suite'] + ')', mensaje=e['mensaje'])
        elif c['id'] in manuales:
            m = manuales[c['id']]
            fila.update(resultado=m['resultado'], fecha=m['fecha'], mensaje=m['mensaje'],
                        fuente='Manual' + (' (' + m['ejecutor'] + ')' if m['ejecutor'] else ''))
        elif c['tipo'] == 'Manual':
            fila['resultado'] = 'Pendiente manual'
        filas.append(fila)
    return datos, filas, ejecuciones


# ---------------------------------------------------------------- Excel
def excel(datos, filas, ejecuciones, destino):
    F = 'Arial'
    base, negrita = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    cab_font, cab_fill = Font(name=F, size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    linea = Side(style='thin', color='BFBFBF')
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    ajuste = Alignment(wrap_text=True, vertical='top')
    centro = Alignment(horizontal='center', vertical='top', wrap_text=True)

    def cabecera(ws, fila, textos, anchos):
        for i, (t, a) in enumerate(zip(textos, anchos), 1):
            c = ws.cell(fila, i, t)
            c.font, c.fill, c.alignment, c.border = cab_font, cab_fill, centro, borde
            ws.column_dimensions[get_column_letter(i)].width = a

    def colorear(ws, rango):
        for r, color in COLOR.items():
            ws.conditional_formatting.add(rango, CellIsRule(operator='equal', formula=[f'"{r}"'],
                                                            fill=PatternFill('solid', fgColor=color)))

    wb = Workbook()

    # ---- Detalle
    det = wb.active
    det.title = 'Detalle'
    cols = ['ID Azure', 'Título', 'HU', 'Responsable', 'Tipo', 'Suite', 'Caso Katalon', 'Resultado',
            'Duración (s)', 'Fecha ejecución', 'Fuente del resultado', 'Mensaje / observación', 'Nota', 'Hallazgos']
    cabecera(det, 1, cols, [9, 48, 34, 24, 15, 20, 44, 14, 11, 17, 22, 50, 36, 12])
    for n, f in enumerate(filas, 2):
        valores = [f['id'], f['titulo'], '\n'.join(f"{h['id']} - {h['titulo']}" for h in f['hu']), f['responsable'],
                   f['tipo'], f['suite'], f['katalon'], f['resultado'], f['duracion'], f['fecha'], f['fuente'],
                   f['mensaje'], f['nota'], f['hallazgos']]
        for i, v in enumerate(valores, 1):
            c = det.cell(n, i, v)
            c.font, c.border = base, borde
            c.alignment = centro if i in (1, 5, 8, 9, 10) else ajuste
        det.cell(n, 1).hyperlink = AZURE + str(f['id'])
        det.cell(n, 1).font = Font(name=F, size=10, color='0563C1', underline='single')
    ultima_det = len(filas) + 1
    det.freeze_panes = 'C2'
    det.auto_filter.ref = f'A1:N{ultima_det}'
    colorear(det, f'H2:H{ultima_det}')

    # ---- HU_Caso (una fila por par HU-caso, para contar por HU)
    huc = wb.create_sheet('HU_Caso')
    cabecera(huc, 1, ['HU', 'Título HU', 'ID caso', 'Resultado'], [8, 60, 10, 16])
    n = 2
    for f in filas:
        for h in f['hu']:
            huc.cell(n, 1, h['id']); huc.cell(n, 2, h['titulo']); huc.cell(n, 3, f['id'])
            huc.cell(n, 4, f'=INDEX(Detalle!$H$2:$H${ultima_det},MATCH(C{n},Detalle!$A$2:$A${ultima_det},0))')
            for i in range(1, 5):
                huc.cell(n, i).font, huc.cell(n, i).border = base, borde
            n += 1
    ultima_huc = n - 1

    # ---- Resumen
    res = wb.create_sheet('Resumen', 0)
    res['A1'] = 'Resultados de pruebas - Conexión360'
    res['A1'].font = Font(name=F, size=14, bold=True)
    res['A2'] = f"Generado: {datetime.now():%Y-%m-%d %H:%M}  ·  Casos: {datos['fuente']} ({datos['fecha_lectura']})  ·  Ejecuciones de Katalon leídas: {len(ejecuciones)}"
    res['A2'].font = Font(name=F, size=9, italic=True)

    cabecera(res, 4, ['Resultado', 'Casos', '% del total'], [22, 10, 12])
    D = f'Detalle!$H$2:$H${ultima_det}'
    for i, r in enumerate(RESULTADOS, 5):
        res.cell(i, 1, r); res.cell(i, 2, f'=COUNTIF({D},A{i})'); res.cell(i, 3, f'=IF($B$11=0,0,B{i}/$B$11)')
    res.cell(11, 1, 'Total'); res.cell(11, 2, '=SUM(B5:B10)'); res.cell(11, 3, '=IF(B11=0,0,B11/B11)')
    res.cell(12, 1, 'Ejecutados'); res.cell(12, 2, '=B5+B6+B7')
    res.cell(13, 1, '% de aprobación (pasó / ejecutados)'); res.cell(13, 2, '=IF(B12=0,0,B5/B12)')
    for r in range(5, 14):
        for c in range(1, 4):
            cel = res.cell(r, c)
            cel.font, cel.border = (negrita if r >= 11 else base), borde
        res.cell(r, 3).number_format = '0%'
    res['B13'].number_format = '0%'
    colorear(res, 'A5:A10')

    cabecera(res, 15, ['Tipo de caso', 'Casos', 'Pasó', 'Falló / Error', 'Sin ejecutar'], [22, 10, 12, 14, 14])
    T = f'Detalle!$E$2:$E${ultima_det}'
    for i, t in enumerate(['Automatizado', 'Automatizado parcial', 'Manual'], 16):
        res.cell(i, 1, t)
        res.cell(i, 2, f'=COUNTIF({T},A{i})')
        res.cell(i, 3, f'=COUNTIFS({T},A{i},{D},"Pasó")')
        res.cell(i, 4, f'=COUNTIFS({T},A{i},{D},"Falló")+COUNTIFS({T},A{i},{D},"Error")')
        res.cell(i, 5, f'=B{i}-C{i}-D{i}')
        for c in range(1, 6):
            res.cell(i, c).font, res.cell(i, c).border = base, borde

    cabecera(res, 20, ['HU', 'Historia de usuario', 'Casos', 'Pasó', 'Falló / Error', 'Sin ejecutar'], [22, 60, 10, 12, 14, 14])
    res.column_dimensions['B'].width = 60
    H, R = f'HU_Caso!$A$2:$A${ultima_huc}', f'HU_Caso!$D$2:$D${ultima_huc}'
    hus = sorted({(h['id'], h['titulo']) for f in filas for h in f['hu']})
    for i, (hid, htit) in enumerate(hus, 21):
        res.cell(i, 1, hid); res.cell(i, 2, htit)
        res.cell(i, 3, f'=COUNTIF({H},A{i})')
        res.cell(i, 4, f'=COUNTIFS({H},A{i},{R},"Pasó")')
        res.cell(i, 5, f'=COUNTIFS({H},A{i},{R},"Falló")+COUNTIFS({H},A{i},{R},"Error")')
        res.cell(i, 6, f'=C{i}-D{i}-E{i}')
        for c in range(1, 7):
            res.cell(i, c).font, res.cell(i, c).border = base, borde
        res.cell(i, 1).alignment = centro
    res.cell(21 + len(hus) + 1, 1, 'Un caso enlazado a varias HU cuenta en cada una de ellas.').font = Font(name=F, size=9, italic=True)
    res.column_dimensions['A'].width = 34

    # ---- Hallazgos (revisión de código y de Azure, no de ejecuciones)
    hz = cargar_hallazgos()
    hal = wb.create_sheet('Hallazgos', 1)
    hal['A1'] = 'Hallazgos de la revisión'
    hal['A1'].font = Font(name=F, size=14, bold=True)
    hal['A2'] = f"{hz['metodo']} ({hz['fecha']})"
    hal['A2'].font = Font(name=F, size=9, italic=True)
    cabecera(hal, 4, ['ID', 'Categoría', 'Severidad', 'Hallazgo', 'Detalle', 'Evidencia', 'Casos de Azure', 'Estado',
                      'Recomendación'], [7, 14, 11, 40, 60, 45, 18, 36, 40])
    sev_color = {'Crítica': 'F4B6B0', 'Alta': 'F8D3A8', 'Media': 'FFF2CC', 'Baja': 'E2EFDA'}
    for n, h in enumerate(hz['hallazgos'], 5):
        for i, v in enumerate([h['id'], h['categoria'], h['severidad'], h['titulo'], h['detalle'], h['evidencia'],
                               ', '.join(map(str, h['casos'])), h['estado'], h['recomendacion']], 1):
            c = hal.cell(n, i, v)
            c.font, c.border = (negrita if i in (1, 4) else base), borde
            c.alignment = centro if i in (1, 2, 3) else ajuste
        if h['severidad'] in sev_color:
            hal.cell(n, 3).fill = PatternFill('solid', fgColor=sev_color[h['severidad']])
    hal.freeze_panes = 'E5'

    # ---- Ejecuciones (historial completo)
    ej = wb.create_sheet('Ejecuciones')
    cabecera(ej, 1, ['Fecha', 'Suite', 'ID Azure', 'Caso Katalon', 'Resultado', 'Duración (s)', 'Mensaje', 'Reporte'],
             [18, 22, 9, 50, 13, 11, 60, 50])
    for n, e in enumerate(reversed(ejecuciones), 2):
        for i, v in enumerate([e['fecha'], e['suite'], e['id'], e['caso'], e['resultado'], e['duracion'],
                               e['mensaje'], e['reporte']], 1):
            c = ej.cell(n, i, v)
            c.font, c.border, c.alignment = base, borde, ajuste
    if not ejecuciones:
        ej.cell(2, 1, 'Todavía no hay ejecuciones de Katalon en Reports/.').font = Font(name=F, size=10, italic=True)
    else:
        colorear(ej, f'E2:E{len(ejecuciones) + 1}')
    ej.freeze_panes = 'A2'

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


# ---------------------------------------------------------------- HTML
def html_reporte(datos, filas, ejecuciones, destino):
    e = lambda s: html.escape(str(s if s is not None else ''))
    total = len(filas)
    conteo = {r: sum(f['resultado'] == r for f in filas) for r in RESULTADOS}
    ejecutados = conteo['Pasó'] + conteo['Falló'] + conteo['Error']
    aprobacion = round(100 * conteo['Pasó'] / ejecutados) if ejecutados else 0
    automatizados = sum(f['tipo'] != 'Manual' for f in filas)
    clase = {'Pasó': 'ok', 'Falló': 'ko', 'Error': 'err', 'Bloqueado': 'blk', 'No ejecutado': 'ne',
             'Pendiente manual': 'pm'}

    barra = ''.join(f'<span class="seg {clase[r]}" style="flex:{conteo[r]}" title="{e(r)}: {conteo[r]}"></span>'
                    for r in RESULTADOS if conteo[r])
    leyenda = ''.join(f'<li><span class="dot {clase[r]}"></span>{e(r)} <b>{conteo[r]}</b></li>' for r in RESULTADOS)

    por_hu = {}
    for f in filas:
        for h in f['hu']:
            d = por_hu.setdefault(h['id'], {'titulo': h['titulo'], 'casos': 0, 'ok': 0, 'ko': 0})
            d['casos'] += 1
            d['ok'] += f['resultado'] == 'Pasó'
            d['ko'] += f['resultado'] in ('Falló', 'Error')
    filas_hu = ''.join(
        f'<tr><td class="num">{hid}</td><td>{e(d["titulo"])}</td><td class="num">{d["casos"]}</td>'
        f'<td class="num">{d["ok"]}</td><td class="num">{d["ko"]}</td><td class="num">{d["casos"] - d["ok"] - d["ko"]}</td>'
        f'<td><div class="mini"><span class="seg ok" style="flex:{d["ok"]}"></span><span class="seg ko" style="flex:{d["ko"]}"></span>'
        f'<span class="seg ne" style="flex:{d["casos"] - d["ok"] - d["ko"]}"></span></div></td></tr>'
        for hid, d in sorted(por_hu.items()))

    filas_det = ''.join(
        f'<tr data-res="{e(f["resultado"])}" data-tipo="{e(f["tipo"])}">'
        f'<td class="num"><a href="{AZURE}{f["id"]}" target="_blank" rel="noopener">{f["id"]}</a></td>'
        f'<td>{e(f["titulo"])}<div class="sub">{e(" · ".join(str(h["id"]) + " " + h["titulo"] for h in f["hu"]))}</div></td>'
        f'<td>{e(f["tipo"])}<div class="sub">{e(f["suite"])}</div></td>'
        f'<td><span class="chip {clase[f["resultado"]]}">{e(f["resultado"])}</span></td>'
        f'<td class="num">{e(f["duracion"] if f["duracion"] is not None else "")}</td>'
        f'<td>{e(f["fecha"])}<div class="sub">{e(f["fuente"])}</div></td>'
        f'<td class="msg">{e(f["mensaje"] or f["nota"])}'
        f'{"".join(f"<span class=hz>{e(x)}</span>" for x in f["hallazgos"].split(", ") if x)}</td></tr>'
        for f in filas)

    hz = cargar_hallazgos()
    sev = {'Crítica': 'crit', 'Alta': 'alta', 'Media': 'media', 'Baja': 'baja'}
    tarjetas_hz = ''.join(
        f'<article class="hzc"><header><span class="sev {sev.get(h["severidad"], "baja")}">{e(h["severidad"])}</span>'
        f'<b>{e(h["id"])}</b> <span class="sub">{e(h["categoria"])}</span></header>'
        f'<h3>{e(h["titulo"])}</h3><p>{e(h["detalle"])}</p>'
        f'<p class="sub"><b>Estado:</b> {e(h["estado"])}</p>'
        f'<p class="sub"><b>Recomendación:</b> {e(h["recomendacion"])}</p>'
        f'<p class="sub"><b>Evidencia:</b> {e(h["evidencia"])}</p>'
        + (f'<p class="sub"><b>Casos de Azure:</b> {e(", ".join(map(str, h["casos"])))}</p>' if h['casos'] else '')
        + '</article>'
        for h in hz['hallazgos'])
    n_crit = sum(h['severidad'] == 'Crítica' for h in hz['hallazgos'])
    aviso = '' if ejecuciones else (
        '<div class="aviso">Todavía no hay ejecuciones de Katalon: los resultados por caso están en "No ejecutado" o '
        '"Pendiente manual". Lo que hay por ahora son los <a href="#hallazgos">hallazgos de la revisión</a>.</div>')

    pagina = f'''<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Resultados de pruebas Conexión360</title>
<style>
:root{{--bg:#f6f7f9;--card:#fff;--tx:#1d2330;--mut:#5f6b7a;--bd:#dde2e8;--ac:#1f4e78;
--ok:#2e8b57;--ko:#c0392b;--err:#d68910;--blk:#7d5ba6;--ne:#9aa5b1;--pm:#d4ac0d}}
@media (prefers-color-scheme:dark){{:root{{--bg:#12161c;--card:#1b2129;--tx:#e6eaf0;--mut:#9aa6b4;--bd:#2c3540;--ac:#6fa8dc}}}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--tx);font:14px/1.45 system-ui,Segoe UI,Arial,sans-serif}}
main{{max-width:1200px;margin:0 auto;padding:24px 16px 48px}}h1{{margin:0 0 4px;font-size:24px}}h2{{font-size:17px;margin:28px 0 10px}}
.meta{{color:var(--mut);font-size:12.5px}}.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:18px 0}}
.card{{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px}}.card b{{display:block;font-size:26px}}.card span{{color:var(--mut);font-size:12.5px}}
.bar,.mini{{display:flex;height:14px;border-radius:7px;overflow:hidden;background:var(--bd)}}.mini{{height:8px;min-width:90px}}
.seg.ok,.dot.ok,.chip.ok{{background:var(--ok)}}.seg.ko,.dot.ko,.chip.ko{{background:var(--ko)}}.seg.err,.dot.err,.chip.err{{background:var(--err)}}
.seg.blk,.dot.blk,.chip.blk{{background:var(--blk)}}.seg.ne,.dot.ne,.chip.ne{{background:var(--ne)}}.seg.pm,.dot.pm,.chip.pm{{background:var(--pm)}}
.legend{{display:flex;flex-wrap:wrap;gap:14px;list-style:none;padding:0;margin:10px 0 0;font-size:13px}}.dot{{display:inline-block;width:10px;height:10px;border-radius:50%;margin-right:6px}}
.panel{{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px;overflow-x:auto}}
table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{text-align:left;padding:7px 8px;border-bottom:1px solid var(--bd);vertical-align:top}}
th{{color:var(--mut);font-weight:600;font-size:12px;position:sticky;top:0;background:var(--card)}}.num{{text-align:right;white-space:nowrap}}
.sub{{color:var(--mut);font-size:11.5px;margin-top:2px}}.msg{{max-width:380px;font-size:12px;color:var(--mut);word-break:break-word}}
.chip{{color:#fff;border-radius:999px;padding:2px 9px;font-size:12px;white-space:nowrap}}a{{color:var(--ac)}}
.filters{{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:10px}}input,select{{font:inherit;padding:6px 8px;border:1px solid var(--bd);border-radius:6px;background:var(--card);color:var(--tx)}}
input{{flex:1;min-width:180px}}
.aviso{{background:var(--card);border:1px solid var(--pm);border-left:5px solid var(--pm);border-radius:8px;padding:10px 14px;margin:14px 0;font-size:13.5px}}
.hzgrid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:12px}}
.hzc{{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px}}.hzc h3{{font-size:14.5px;margin:8px 0 6px}}.hzc p{{margin:6px 0;font-size:13px}}
.hzc header{{display:flex;align-items:center;gap:8px}}.sev{{color:#fff;border-radius:999px;padding:2px 9px;font-size:11.5px;font-weight:600}}
.sev.crit{{background:#b03a2e}}.sev.alta{{background:#d35400}}.sev.media{{background:#b7950b}}.sev.baja{{background:#5d6d7e}}
.hz{{display:inline-block;margin:4px 4px 0 0;padding:1px 7px;border:1px solid var(--bd);border-radius:6px;font-size:11px;color:var(--tx)}}
</style></head><body><main>
<h1>Resultados de pruebas · Conexión360</h1>
<div class="meta">Generado {datetime.now():%Y-%m-%d %H:%M} · Casos: {e(datos['fuente'])} ({e(datos['fecha_lectura'])}) · Ejecuciones de Katalon leídas: {len(ejecuciones)}</div>
<section class="cards">
<div class="card"><b>{total}</b><span>Test cases en Azure</span></div>
<div class="card"><b>{automatizados}</b><span>Automatizados en Katalon</span></div>
<div class="card"><b>{ejecutados}</b><span>Ejecutados</span></div>
<div class="card"><b>{conteo['Pasó']}</b><span>Pasaron</span></div>
<div class="card"><b>{conteo['Falló'] + conteo['Error']}</b><span>Fallaron / error</span></div>
<div class="card"><b>{aprobacion}%</b><span>Aprobación (pasó / ejecutados)</span></div>
</section>
{aviso}
<div class="panel"><div class="bar">{barra}</div><ul class="legend">{leyenda}</ul></div>
<h2 id="hallazgos">Hallazgos de la revisión ({len(hz['hallazgos'])}, {n_crit} críticos)</h2>
<div class="meta" style="margin-bottom:10px">{e(hz['metodo'])} ({e(hz['fecha'])})</div>
<div class="hzgrid">{tarjetas_hz}</div>
<h2>Por historia de usuario</h2>
<div class="panel"><table><thead><tr><th class="num">HU</th><th>Historia de usuario</th><th class="num">Casos</th><th class="num">Pasó</th><th class="num">Falló</th><th class="num">Sin ejecutar</th><th></th></tr></thead>
<tbody>{filas_hu}</tbody></table><div class="sub">Un caso enlazado a varias HU cuenta en cada una de ellas.</div></div>
<h2>Detalle por test case</h2>
<div class="panel">
<div class="filters"><input id="q" type="search" placeholder="Buscar por ID, título o HU…" aria-label="Buscar">
<select id="fr" aria-label="Resultado"><option value="">Todos los resultados</option>{''.join(f'<option>{e(r)}</option>' for r in RESULTADOS)}</select>
<select id="ft" aria-label="Tipo"><option value="">Todos los tipos</option><option>Automatizado</option><option>Automatizado parcial</option><option>Manual</option></select></div>
<table id="t"><thead><tr><th class="num">ID</th><th>Test case / HU</th><th>Tipo</th><th>Resultado</th><th class="num">Duración (s)</th><th>Ejecución</th><th>Mensaje / nota</th></tr></thead>
<tbody>{filas_det}</tbody></table></div>
</main>
<script>
const q=document.getElementById('q'),fr=document.getElementById('fr'),ft=document.getElementById('ft');
function filtrar(){{const s=q.value.toLowerCase();document.querySelectorAll('#t tbody tr').forEach(tr=>{{
tr.style.display=(!s||tr.innerText.toLowerCase().includes(s))&&(!fr.value||tr.dataset.res===fr.value)&&(!ft.value||tr.dataset.tipo===ft.value)?'':'none';}});}}
[q,fr,ft].forEach(x=>x.addEventListener('input',filtrar));
</script></body></html>'''
    destino.write_text(pagina, encoding='utf-8')


if __name__ == '__main__':
    datos, filas, ejecuciones = consolidar()
    sello = datetime.now().strftime('%Y%m%d_%H%M')
    xlsx, htm = AQUI / f'Reporte_Resultados_{sello}.xlsx', AQUI / f'Reporte_Resultados_{sello}.html'
    excel(datos, filas, ejecuciones, xlsx)
    html_reporte(datos, filas, ejecuciones, htm)
    print(f'{len(filas)} casos, {len(ejecuciones)} ejecuciones de Katalon')
    print(xlsx.name)
    print(htm.name)
